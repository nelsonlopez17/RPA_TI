"""
REPORTE DE TOP 10 PRODUCTOS

Genera un reporte analítico de los últimos 7 días. Consulta las facturas y sus detalles, 
calcula cuáles son los 10 productos más vendidos (por cantidad) y exporta 
los resultados automáticamente a un archivo Excel formateado.
"""
import sys
import time
import datetime
from collections import defaultdict

try:
    import requests
except ImportError:
    import urllib.request as _urllib
    import json as _json
    class SimpleRequests:
        @staticmethod
        def get(url):
            resp = _urllib.urlopen(url)
            return type('Response', (), {'status_code': resp.getcode(), 'json': lambda self=None: _json.loads(resp.read().decode())})()
    requests = SimpleRequests

# Forzar codificación UTF-8 en la consola de Windows para evitar errores con emojis
sys.stdout.reconfigure(encoding='utf-8')

BASE_URL = "https://erp-django-smea.onrender.com/api/v1"
DETALLES_URL = f"{BASE_URL}/detalle-facturas/"
FACTURAS_URL = f"{BASE_URL}/facturas/"


def get_all_pages(url):
    """Obtiene todos los registros paginados de un endpoint."""
    resultados = []
    next_url = url
    while next_url:
        try:
            resp = requests.get(next_url, timeout=30)
        except Exception as e:
            print(f"Error de conexion con el servidor: {e}")
            break
        if resp.status_code == 404:
            print(f"\nERROR: El endpoint no fue encontrado en el servidor.")
            print(f"URL: {next_url}")
            print(f"SOLUCION: Asegurate de haber hecho 'git push' y que Render haya desplegado los cambios.")
            break
        if resp.status_code != 200:
            print(f"Error al consultar {next_url}: HTTP {resp.status_code}")
            break
        data = resp.json()
        if isinstance(data, list):
            resultados.extend(data)
            break
        else:
            resultados.extend(data.get('results', []))
            next_url = data.get('next')
    return resultados


def run_bot():
    print("==============================================")
    print("  TOP 10 PRODUCTOS MAS FACTURADOS - ERP BOT  ")
    print("==============================================")
    print(f"Consultando datos en: {BASE_URL}")
    print("Obteniendo facturas para filtrar por los ultimos 7 dias...")
    try:
        facturas = get_all_pages(FACTURAS_URL)
        
        # Calcular fecha limite (hace 7 dias)
        fecha_limite = datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(days=7)
        facturas_recientes_ids = set()
        factura_fechas_map = {}
        
        for f in facturas:
            fecha_str = f.get('fecha_venta')
            if fecha_str:
                try:
                    fecha_str = fecha_str.replace("Z", "+00:00")
                    fecha_venta = datetime.datetime.fromisoformat(fecha_str)
                    if fecha_venta >= fecha_limite:
                        f_id = f.get('id')
                        facturas_recientes_ids.add(f_id)
                        factura_fechas_map[f_id] = fecha_venta.strftime("%Y-%m-%d")
                except ValueError:
                    pass

        print("Obteniendo detalles de facturas... (puede tardar unos segundos)")
        detalles = get_all_pages(DETALLES_URL)

        if not detalles:
            print("No se encontraron detalles de facturas en el sistema.")
            return
            
        detalles_recientes = [d for d in detalles if d.get('factura') in facturas_recientes_ids]
        
        if not detalles_recientes:
            print("No se encontraron detalles de ventas en los ultimos 7 dias.")
            return
            
        print(f"Se procesaran {len(detalles_recientes)} detalles de {len(facturas_recientes_ids)} facturas recientes.\n")

        # Agrupar por producto
        productos_stats = defaultdict(lambda: {
            'nombre': '',
            'cantidad_total': 0,
            'monto_total': 0.0,
            'clientes': set(),
            'fechas': set()
        })

        for d in detalles_recientes:
            prod_id = d.get('producto')
            prod_nombre = d.get('producto_nombre', f'Producto #{prod_id}')
            cantidad = d.get('cantidad', 0)
            precio_unit = float(d.get('precio_unitario', 0))
            cliente = d.get('cliente_nombre', 'Desconocido')
            f_id = d.get('factura')

            productos_stats[prod_id]['nombre'] = prod_nombre
            productos_stats[prod_id]['cantidad_total'] += cantidad
            productos_stats[prod_id]['monto_total'] += cantidad * precio_unit
            productos_stats[prod_id]['clientes'].add(cliente)
            if f_id in factura_fechas_map:
                productos_stats[prod_id]['fechas'].add(factura_fechas_map[f_id])

        # Ordenar por cantidad total descendente y tomar top 10
        top_10 = sorted(
            productos_stats.items(),
            key=lambda x: x[1]['cantidad_total'],
            reverse=True
        )[:10]

        # Mostrar en consola
        print(f"{'#':<4} {'Producto':<30} {'Fechas':<20} {'Clientes':<30} {'Cant.':<8} {'Monto Total'}")
        print("-" * 110)
        for rank, (prod_id, stats) in enumerate(top_10, 1):
            clientes_str = ", ".join(sorted(stats['clientes']))
            fechas_str = ", ".join(sorted(stats['fechas']))
            # Truncar si es muy largo para la consola
            prod_display = stats['nombre'][:28] + ".." if len(stats['nombre']) > 30 else stats['nombre']
            cli_display = clientes_str[:28] + ".." if len(clientes_str) > 30 else clientes_str
            fechas_display = fechas_str[:18] + ".." if len(fechas_str) > 20 else fechas_str
            print(f"{rank:<4} {prod_display:<30} {fechas_display:<20} {cli_display:<30} {stats['cantidad_total']:<8} Q.{stats['monto_total']:,.2f}")

        print("-" * 110)
        print(f"\nTotal de productos analizados: {len(productos_stats)}")

        # Exportar a Excel
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

            excel_path = "top_10_productos.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Top 10 Productos"

            # Titulo principal
            ws.merge_cells('A1:F1')
            titulo = ws['A1']
            titulo.value = "TOP 10 PRODUCTOS MAS VENDIDOS - ERP"
            titulo.font = Font(bold=True, size=14, color="FFFFFF")
            titulo.alignment = Alignment(horizontal="center", vertical="center")
            titulo.fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            ws.row_dimensions[1].height = 30

            # Encabezados de columna
            header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            header_align = Alignment(horizontal="center", vertical="center")
            headers = ['Ranking', 'Producto', 'Fechas', 'Clientes', 'Cantidad Total', 'Monto Total (Q)']
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align

            # Datos
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )
            for rank, (prod_id, stats) in enumerate(top_10, 1):
                clientes_str = ", ".join(sorted(stats['clientes']))
                fechas_str = ", ".join(sorted(stats['fechas']))
                row = rank + 2
                ws.cell(row=row, column=1, value=rank).alignment = Alignment(horizontal="center")
                ws.cell(row=row, column=2, value=stats['nombre'])
                ws.cell(row=row, column=3, value=fechas_str)
                ws.cell(row=row, column=4, value=clientes_str)
                ws.cell(row=row, column=5, value=stats['cantidad_total']).alignment = Alignment(horizontal="center")
                monto = ws.cell(row=row, column=6, value=round(stats['monto_total'], 2))
                monto.number_format = '"Q."#,##0.00'
                # Fondo alterno para legibilidad
                if rank % 2 == 0:
                    alt_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
                    for col in range(1, 7):
                        ws.cell(row=row, column=col).fill = alt_fill
                # Bordes
                for col in range(1, 7):
                    ws.cell(row=row, column=col).border = thin_border

            # Ajustar anchos de columna
            ws.column_dimensions['A'].width = 10
            ws.column_dimensions['B'].width = 35
            ws.column_dimensions['C'].width = 25
            ws.column_dimensions['D'].width = 45
            ws.column_dimensions['E'].width = 16
            ws.column_dimensions['F'].width = 20

            wb.save(excel_path)
            print(f"\nReporte exportado exitosamente: {excel_path}")

        except ImportError:
            print("\nAVISO: No se pudo exportar a Excel. Instala openpyxl con: pip install openpyxl")

    except Exception as e:
        print(f"Error critico durante la ejecucion: {e}")


if __name__ == "__main__":
    run_bot()
