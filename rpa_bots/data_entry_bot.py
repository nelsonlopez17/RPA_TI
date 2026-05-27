"""
BOT DE DATA ENTRY (REGISTRO MASIVO)

Simula a un operador humano registrando grandes volúmenes de inventario.
Lee el archivo Excel 'data/nuevos_productos.xlsx', procesa la información y 
se conecta a la API del ERP para inyectar los productos en la base de datos central.
"""
import requests
import time
import sys

# Forzar codificación UTF-8 en la consola de Windows para evitar errores con emojis
sys.stdout.reconfigure(encoding='utf-8')


API_URL = "https://erp-django-smea.onrender.com/api/v1/productos/"
EXCEL_PATH = "data/nuevos_productos.xlsx"

def run_bot():
    print("🤖 Iniciando Bot de Data Entry...")
    print(f"📄 Leyendo archivo: {EXCEL_PATH}")
    
    try:
        import openpyxl
    except ImportError:
        print("❌ Error: No se encontró el módulo 'openpyxl'. Instálalo con 'pip install openpyxl'")
        return

    try:
        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active
        
        headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
                
            row_dict = dict(zip(headers, row))
            nombre = str(row_dict.get('nombre') or '').strip()
            sku = str(row_dict.get('sku') or '').strip()
            
            if not nombre and not sku:
                continue
                
            if not nombre or not sku:
                print("⚠️ Fila ignorada: Faltan datos obligatorios (nombre o sku).")
                continue
                
            print(f"⏳ Procesando: {nombre} (SKU: {sku})")
            
            precio_compra = float(row_dict.get('precio_compra') or 0.0)
            precio_venta = round(precio_compra * 1.25, 2)
            
            payload = {
                "nombre": nombre,
                "sku": sku,
                "categoria": int(row_dict.get('categoria') or 1),
                "descripcion": str(row_dict.get('descripcion') or ''),
                "precio_compra": precio_compra,
                "precio_venta": precio_venta,
                "stock": int(row_dict.get('stock') or 0)
            }
            
            try:
                response = requests.post(API_URL, json=payload)
                if response.status_code == 201:
                    print(f"✅ ¡Éxito! Producto '{nombre}' registrado en el ERP.")
                else:
                    print(f"❌ Error al registrar '{nombre}': {response.status_code} - {response.text}")
            except Exception as e:
                print(f"❌ Error de conexión: {e}")
            
            # Simular tiempo de carga humano
            time.sleep(1)
            
    except FileNotFoundError:
        print(f"❌ Error: No se encontró el archivo '{EXCEL_PATH}'. Puedes usar la plantilla generada.")
    except Exception as e:
        print(f"❌ Error inesperado: {e}")
            
    print("🏁 Bot de Data Entry finalizó sus tareas.")

if __name__ == "__main__":
    run_bot()
